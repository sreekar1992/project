"""Build the evidence-based IEEE-style manuscript using the bundled docx runtime."""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/ieee_paper'
FIG = OUT / 'figures'
RESULTS = json.loads((ROOT / 'output/paper_results/results.json').read_text())
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.75), Inches(1)
section.left_margin = section.right_margin = Inches(.625)
section.header_distance = section.footer_distance = Inches(.3)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Times New Roman', Pt(10)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.first_line_indent = Inches(.14)
normal.paragraph_format.widow_control = True
for name in ['Title', 'Heading 1', 'Heading 2', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.first_line_indent = 0
    st.paragraph_format.line_spacing = 1
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.bold = False
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 1'].font.size = Pt(10)
doc.styles['Heading 1'].font.bold = False
doc.styles['Heading 1'].font.small_caps = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Heading 2'].font.size = Pt(10)
doc.styles['Heading 2'].font.bold = False
doc.styles['Heading 2'].font.italic = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
doc.styles['Caption'].font.size = Pt(8)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].font.bold = False
doc.styles['Caption'].paragraph_format.space_after = Pt(7)
# Remove built-in theme fonts and title rules; IEEE-style typography is explicit.
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption']:
    rpr = doc.styles[name].element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    for key in list(fonts.attrib):
        if key.endswith('Theme'):
            del fonts.attrib[key]
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), 'Times New Roman')
doc.core_properties.title = 'Adversarial ECG Camouflage with Authenticated Recovery and Residual Attention Classification'
doc.core_properties.author = 'Muddisetty Sreekar varma; Sachikanta dash'
doc.core_properties.subject = 'ECG signal classification, adversarial sensitivity, and exact authenticated recovery'
doc.core_properties.keywords = 'ECG, adversarial camouflage, authenticated encryption, IEEE style, 26MPE05015'


def p(text, bold=False, indent=True, size=10):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if not indent:
        para.paragraph_format.first_line_indent = 0
    for item in re.split(r'(?:\b)(epsilon_[ac]|Epsilon_[ac]|U\(g_t\))', text):
        if re.fullmatch(r'[Ee]psilon_[ac]', item):
            obj = OxmlElement('m:oMath'); obj.append(sub('ε', item[-1])); para._p.append(obj)
        elif item == 'U(g_t)':
            obj = OxmlElement('m:oMath'); obj.append(mathrun('U(')); obj.append(sub('g', 't')); obj.append(mathrun(')')); para._p.append(obj)
        else:
            run = para.add_run(item)
            run.bold, run.font.size = bold, Pt(size)
    return para


def h(text, level=1):
    return doc.add_paragraph(text, f'Heading {level}')


def table(caption, headers, rows, widths):
    cap = doc.add_paragraph(caption, 'Caption')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(4)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            el.set(qn('w:' + key), value)
        borders.append(el)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ['top', 'left', 'bottom', 'right']:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:w'), '45' if edge in ['top', 'bottom'] else '65')
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    props.append(margins)
    for i, label in enumerate(headers):
        t.rows[0].cells[i].text = label
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EEEEEE')
                cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.first_line_indent = 0
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(8)
                    run.font.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def figure(filename, caption):
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = 0
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.space_after = Pt(2)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic = para.add_run().add_picture(str(FIG / filename), width=Inches(3.45))
    pic._inline.docPr.set('descr', caption)
    cap = doc.add_paragraph(caption, 'Caption')
    cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def mathrun(text):
    run = OxmlElement('m:r')
    rp = OxmlElement('w:rPr')
    fs = OxmlElement('w:sz'); fs.set(qn('w:val'), '18'); rp.append(fs)
    run.append(rp)
    value = OxmlElement('m:t'); value.text = text; run.append(value)
    return run


def sub(base, index):
    node = OxmlElement('m:sSub')
    for tag, value in [('e', base), ('sub', index)]:
        el = OxmlElement('m:' + tag); el.append(mathrun(value)); node.append(el)
    return node


def equation(parts, number):
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = 0
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(6)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    obj = OxmlElement('m:oMath')
    for item in parts:
        obj.append(mathrun(item) if isinstance(item, str) else sub(*item))
    para._p.append(obj)
    para.add_run('   (' + str(number) + ')').font.size = Pt(9)


title = doc.add_paragraph(doc.core_properties.title, 'Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
authors = doc.add_paragraph('Muddisetty Sreekar varma and Sachikanta dash')
authors.paragraph_format.first_line_indent = 0
authors.paragraph_format.space_after = Pt(12)
authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
authors.runs[0].font.size = Pt(12)
body_section = doc.add_section(WD_SECTION_START.CONTINUOUS)
cols = body_section._sectPr.find(qn('w:cols'))
cols.set(qn('w:num'), '2'); cols.set(qn('w:space'), '360')

p('Abstract—We present a local research workbench that combines electrocardiogram rhythm classification, adversarial signal camouflage, and authenticated recovery of original MATLAB recordings. A compact one-dimensional residual-attention network processes 1,000 ten-second MLII fragments from a public 17-class dataset. The reconstructed 800/200 fragment split gives 87.00% accuracy and 86.22% weighted F1; all validation recording identifiers overlap training and three validation waveforms are exact training duplicates. These figures therefore describe leakage-affected internal validation, not unseen-patient performance. A gradient-guided routine changes a publicly readable waveform while an AES-256-GCM payload retains the complete original file. Authenticated creation-time reports show changed predictions for all 1,000 fragments under one known checkpoint, and a separate cryptographic recheck confirms exact MAT recovery for every fragment. An independent normalized-input FGSM experiment reduces validation accuracy to 59.50% at epsilon 0.05. The study separates classifier sensitivity from cryptographic confidentiality: decoy misclassification is neither attack prevention nor anonymization. The implementation provides a reproducible demonstration of altered-signal prediction and lossless original recovery, without claiming the PPO-based defenses of the RNAF reference framework.', bold=True, indent=False, size=9)
p('Index Terms—Electrocardiography, adversarial examples, authenticated encryption, residual attention, signal classification, reproducibility.', bold=True, indent=False, size=9)

h('I. Introduction')
p('ECG research applications often display a waveform, a predicted rhythm, and a model score on the same screen. When adversarial modification and encryption are added, these representations can be confused. A waveform that receives a different label is still numerical signal data; an encrypted file is a byte sequence that should be authenticated and decrypted before ordinary inference. This distinction is central to the present workbench.')
p('We investigate this distinction using the public ECG signals dataset of Plawiak [1], derived from the MIT-BIH Arrhythmia Database [2]. The application reads MATLAB files directly, predicts one of 17 dataset labels, and visualizes the original, public camouflaged, and restored signals. The source files remain unchanged. A password unlocks the protected original rather than attempting to reverse the adversarial perturbation.')
p('The motivation comes from RNAF, which combines image-based adversarial analysis with reinforcement-learning defenses [3]. Its fish-to-dog illustration demonstrates changed model behavior under input manipulation. Our implementation adapts that experimental idea to one-dimensional ECG signals, but does not implement RNAF\'s PPO policy, adaptive steganalysis, or CIFAR-10 image pipeline. We use the reference to motivate a clearly scoped experiment, not to claim equivalent architecture or results.')
p('The contributions are an integrated three-stage ECG comparison, a raw-waveform camouflage heuristic evaluated through the actual preprocessing chain, and an authenticated container that restores exact original MAT bytes. We also report a dataset-leakage audit and distinguish whole-dataset camouflage outcomes from the baseline validation experiment. The objective is a transparent research prototype, not a deployable clinical diagnosis or privacy system.')

h('II. Related Work')
p('Goodfellow et al. introduced the fast gradient sign method (FGSM), which uses an input-loss gradient to construct adversarial examples [4]. Madry et al. developed a robust-optimization perspective and stronger iterative attacks [5]. These works motivate sensitivity testing, but changing a classifier output is not the same operation as encrypting a recording. A readable perturbed waveform does not acquire a confidentiality guarantee merely because its predicted class changes.')
p('RNAF studies digital images using a combination of adversarial training, PPO-based adaptive defense, and steganalysis [3]. Its task, data, and threat model differ from our ECG demonstration. We do not directly compare their reported accuracy with our 17-class results, because such a comparison would mix datasets, splits, models, and attacks. The present camouflage update is also not exact projected gradient descent through the raw preprocessing chain.')
table('TABLE I\nDistinct roles of the implemented components', ['Component', 'Role', 'Does not establish'], [
    ['Residual-attention CNN', '17-class prediction', 'Clinical diagnosis'],
    ['Raw camouflage', 'Public label-change experiment', 'Confidentiality'],
    ['FGSM test', 'Input sensitivity measurement', 'Universal robustness'],
    ['AES-GCM recovery', 'Protected exact-original storage', 'Safety of public decoy'],
    ['RNAF reference [3]', 'Image-defense motivation', 'Implemented PPO defense'],
], [1.02, 1.18, 1.25])
p('For file protection, we use AES-GCM, an authenticated-encryption mode specified by NIST [6]. For model visualization, the workbench can generate one-dimensional Grad-CAM relevance maps based on gradients and feature activations [7]. Those maps indicate model-dependent relevance, not validated clinical landmarks. Table I summarizes the roles of these components without implying that they provide interchangeable security properties.')

h('III. Dataset and Classification Baseline')
h('A. Data representation and class support', 2)
p('The Mendeley Data version 3 release contains 1,000 ten-second MLII fragments sampled at 360 Hz, with 3,600 samples per file and 17 rhythm or pattern categories [1]. The release describes recordings from 45 patients and is distributed under CC BY 4.0. Our local loader reads the numeric val array and derives the reference class from its parent folder. These labels are not 17 independent cardiovascular disease diagnoses. Table II reports the actual local class counts and reconstructed validation support.')
table('TABLE II\nLocal dataset support and validation recall', ['ID and label', 'All', 'Val.', 'Recall %'], [
    [c['class'], c['dataset_n'], c['validation_n'], f"{100*c['recall']:.1f}"] for c in RESULTS['classes']
], [1.64, .53, .53, .75])
p('Abbreviations follow the dataset folders: NSR, normal sinus rhythm; APB, atrial premature beat; AFL, atrial flutter; AFIB, atrial fibrillation; SVTA, supraventricular tachyarrhythmia; WPW, Wolff-Parkinson-White pattern; PVC, premature ventricular contraction; VT, ventricular tachycardia; IVR, idioventricular rhythm; VFL, ventricular flutter; LBBBB/RBBBB, left/right bundle branch block beat; SDHB, second-degree heart block; and PR, pacemaker rhythm. Fusion denotes fusion of ventricular and normal beats.', size=9)
h('B. Preprocessing and network architecture', 2)
p('The preprocessing function applies a fourth-order 0.5-45 Hz Butterworth band-pass with forward-backward filtering. Each fragment is centered by its median and scaled by 1.4826 times its median absolute deviation, using a floor of 10⁻⁶. Every fourth sample is retained, producing 900 normalized input values at 90 Hz. The classifier operates on these values, not on the plotted PNG. The cutoff lies at the downsampled Nyquist frequency; the current implementation has no separate resampling filter, which is a limitation for future signal-processing refinement.')
p('The custom RAMNV2 model contains 27,633 trainable parameters. It uses a convolutional stem, depthwise/pointwise residual blocks, channel attention, temporal average pooling, and a 17-output linear layer. It is inspired by efficient mobile-network principles rather than being a standard MobileNetV2 image backbone. Within each residual block, a kernel-five depthwise convolution is followed by normalization and ReLU, a pointwise convolution, normalization, channel weighting, and a residual addition. The attention gate pools across time, reduces channels, applies ReLU, expands channels, and applies a sigmoid.')
table('TABLE III\nClassifier stages for one input fragment', ['Stage', 'Channels', 'Length'], [
    ['Normalized input', '1', '900'],
    ['Kernel-7 stem, stride 2', '32', '450'],
    ['Residual blocks, strides 1 and 2', '32', '225'],
    ['Pointwise channel expansion', '64', '225'],
    ['Residual blocks, strides 1, 2, 1', '64', '113'],
    ['Temporal average pooling', '64', '1'],
    ['Linear logits', '17', '1'],
], [2.2, .65, .6])
h('C. Training provenance and split audit', 2)
p('The recorded training command used 30 epochs and batch size 32. The implementation uses cross-entropy and AdamW with learning rate 0.001 and weight decay 0.0001. The legacy checkpoint stores the state dictionary, class count, and label map, but does not retain its training configuration. Seed 42 and the stratified 80/20 split are reconstructed from the original defaults; they reproduce the saved accuracy and weighted metrics exactly. This provenance distinction prevents later metadata-saving code from being retrospectively attributed to the earlier checkpoint.')
p('The reconstruction yields 800 training and 200 validation fragments. All 44 filename-derived validation recording identifiers also occur in training, and all 200 validation fragments consequently belong to overlapping recordings. Three validation waveforms exactly match training waveforms. Recording identifiers are proxies rather than independently verified patient identities. The reported validation scores therefore cannot estimate generalization to unseen recordings or patients.')

h('IV. Camouflage and Authenticated Recovery')
h('A. System workflow and scope', 2)
p('Fig. 1 shows two branches from an original MAT file. One branch creates a public altered waveform. The other encrypts the complete original file bytes. The branches are packaged together, while original source files and model weights remain unchanged. The public waveform can be plotted and classified before password entry. The recovery branch is used only after password-based key unwrapping and successful authenticated decryption.')
figure('fig01_workflow.png', 'Fig. 1. Public-signal and protected-original paths. Recovery reads the encrypted original payload, not an inverse perturbation. The batch key envelope and authenticated manifest are also required.')
p('The white-box camouflage routine knows the selected classifier and its preprocessing. It tries to change that classifier\'s original prediction within a raw-amplitude budget. We do not assume that the resulting waveform conceals identity, prevents an attacker from learning clinical information, or transfers to another model. An attacker who can access the plaintext source directory is outside the protection offered by encrypted copies. The prototype runs locally for a trusted user and is not a remotely hardened service.')
h('B. Iterative raw waveform camouflage', 2)
p('Let x denote the 3,600-sample raw waveform, P the fixed preprocessing operation, and f the classifier. Let s be the scaled median absolute deviation of the band-passed original signal. The raw perturbation budget b and nominal update size alpha are defined by (1), where epsilon_c is the camouflage setting and T is the maximum iteration count. Defaults are epsilon_c = 0.5 and T = 20.')
equation(['b = ', ('ε', 'c'), ' s,    α = 2b/T'], 1)
p('At each iteration the routine computes cross-entropy relative to the original predicted class, differentiates in the 900-sample normalized model-input domain, and linearly interpolates that gradient to 3,600 points. Denote the interpolated gradient by U(g_t). The raw candidate update is (2), where the projection clips each sample to the interval from x minus b to x plus b.')
equation([('x', 't+1'), ' = ', ('Π', 'b'), '(', ('x', 't'), ' + α sign(U(', ('g', 't'), ')))'], 2)
p('Every candidate is run again through the real SciPy preprocessing and classifier. The process stops when the predicted label changes; otherwise the candidate with the lowest original predicted-class softmax score is retained. Bounds are adjusted inward for float32 representation so that rounding does not enlarge the budget. Constant or numerically zero-scale inputs are not modified. Because the method does not differentiate through filtering and median/MAD normalization, it is a gradient-guided heuristic rather than exact raw-space PGD.')
p('Algorithm 1 summarizes the implementation. (1) Read and preserve the original MAT bytes. (2) Measure the original model output and raw budget. (3) Generate projected raw candidates, re-evaluating each candidate through P and f. (4) Record the best actual output and whether its label changed. (5) Encrypt the exact original bytes and bind the public decoy and record context as associated data. (6) Verify exact recovery before publishing the completed batch. A reference label is used for correctness reporting, not to fabricate a changed prediction.', indent=False, size=9)
h('C. File container and key handling', 2)
p('Each batch receives a random 256-bit AES data key. Original MAT files, rendered waveform PNGs, and the private manifest are encrypted using AES-GCM with fresh 96-bit nonces and 128-bit authentication tags. A password-derived wrapping key protects the batch key. The implementation uses scrypt with N = 131072, r = 8, p = 1, and a random 128-bit salt. The demo intentionally uses the public fixed password 987654321; this configuration is unsuitable for confidential patient data despite the standard encryption primitive.')
p('A camouflaged MAT contains public float32 val samples and an encrypted original-MAT recovery payload, together with a nonce, format version, and batch/record context. Associated data bind the version, context, and canonical little-endian float32 decoy bytes. Altering those bound fields causes authentication failure. Recovery retrieves the original file as bytes, preserving its original MATLAB representation rather than reconstructing it from a rounded waveform array. A complete batch archive is required because an individual container does not include a usable unwrapped key.')
p('The GUI checks MAT and PNG equality and compares source and recovered predictions using the same selected checkpoint. Results and explanation images are encrypted when saved. Before password entry, a grayscale diagnostic can visualize ciphertext byte values; this is not an ECG image input or an authenticity check. Public decoy metadata are similarly unauthenticated until recovery succeeds. Clearing results removes decrypted previews from the interface but does not guarantee secure erasure from browser or operating-system memory.')
h('D. Separate FGSM sensitivity experiment', 2)
p('The optional FGSM experiment uses a temporary copy of the normalized 900-sample input. It adds epsilon_a times the sign of the cross-entropy gradient with respect to the reference class [4]. No arbitrary amplitude clipping is applied. Epsilon_a = 0.05 therefore has different units and a different optimization procedure from epsilon_c = 0.5. Neither value is a decryption parameter. Correct restoration depends on the key and authenticated payload, not on matching perturbation settings.')

h('V. Experimental Results')
h('A. Internal classification evaluation', 2)
p('The existing checkpoint correctly classifies 174 of the 200 reconstructed validation fragments. Table IV provides overall metrics, and Fig. 2 shows the confusion matrix. Weighted F1 is 86.22%, whereas macro F1 is 76.09%, indicating that the aggregate weighted score masks uneven class behavior. SVTA and Fusion have zero recall. Several other classes have only two to four validation samples, so apparently perfect scores should not be interpreted as reliable estimates for those rhythms.')
table('TABLE IV\nLeakage-affected internal validation results', ['Metric', 'Value'], [
    ['Correct fragments', '174 / 200'], ['Accuracy', '87.00%'],
    ['Weighted precision', '86.40%'], ['Weighted recall', '87.00%'],
    ['Weighted F1', '86.22%'], ['Macro F1', '76.09%'],
    ['Balanced accuracy', '76.79%'], ['Macro OvR ROC AUC', '0.9907'],
    ['Macro average precision', '0.9151'],
], [2.3, 1.15])
p('The macro one-versus-rest ROC AUC of 0.9907 describes ranking of continuous scores, not the accuracy of the maximum-score label. High ranking scores can coexist with low class recall. Model outputs have not been calibrated as clinical probabilities, and no prospective clinical comparison was performed. We do not plot a training/validation learning curve because per-epoch validation history was not saved.')
h('B. FGSM sensitivity of the baseline', 2)
p('Fig. 3 and Table V summarize the separate white-box FGSM experiment on the same 200 validation fragments. Accuracy decreases from 87.0% without perturbation to 59.5% at epsilon_a = 0.05 and 14.0% at 0.20. The 0.05 setting produces a 27.5-percentage-point absolute decrease. The evaluated model is the original baseline; optional adversarial-training code is available, but those training options are not evidence of an adversarially trained checkpoint in this experiment.')
figure('fig02_confusion_matrix.png', 'Fig. 2. Confusion matrix for 200 validation fragments; class IDs follow Table II. Counts sum to 200, with 174 correct. Recording overlap and exact duplicates limit interpretation.')
figure('fig03_fgsm_curve.png', 'Fig. 3. Baseline FGSM sensitivity on the reconstructed validation split. Epsilon is measured in normalized input units. This single-attack curve is not a robustness guarantee.')
table('TABLE V\nNormalized-input FGSM experiment', ['Epsilon', 'Correct / 200', 'Accuracy %', 'Macro F1'], [
    [f"{r['epsilon']:.2f}", r['correct'], f"{100*r['accuracy']:.1f}", f"{r['macro_f1']:.4f}"] for r in RESULTS['robustness']
], [.55, 1.05, .95, .9])
h('C. Whole-dataset camouflage and exact recovery', 2)
p('A completed 1,000-record camouflage batch uses epsilon_c = 0.5, at most 20 steps, and the same checkpoint. Its authenticated creation-time reports record label changes for all 1,000 fragments. This is a whole-dataset, model-specific optimization experiment that includes training fragments; it is not an independent attack-prevention rate. Of the original predictions, 969 were correct. All 969 became incorrect after camouflage, while 21 initially incorrect predictions became correct. A label-change count therefore differs from a correctness-based attack-success statistic.')
table('TABLE VI\nWhole-dataset camouflage and recovery audit', ['Quantity', 'Observed result'], [
    ['Fragments processed', '1,000'],
    ['Changed predicted labels', '1,000 / 1,000'],
    ['Original correct → incorrect', '969'],
    ['Original incorrect → correct', '21'],
    ['Iterations used min / median / max', '1 / 2 / 9'],
    ['Measured raw-budget violations', '0'],
    ['Exact embedded MAT recovery', '1,000 / 1,000'],
    ['Recovered MAT matches source', '1,000 / 1,000'],
], [2.25, 1.2])
p('Dataset-wide creation-time prediction reports were authenticated without recomputing their predictions; the single Fig. 4 example was re-evaluated. Independently, all embedded recovery payloads were decrypted and compared against the separately encrypted original MAT bytes and current source files. All 1,000 comparisons were exact, and every original MAT and PNG plaintext hash matched the authenticated manifest. This establishes byte-preserving recovery for the audited batch; it does not establish that the readable decoys protect privacy.')
p('For Fig. 4, the decoy changes the selected model\'s prediction from NSR to APB after two iterations. The raw budget is 8.85518 dataset units, the measured maximum change is 1.77100, and decoy-versus-original MSE is 2.37149. Original-versus-restored MSE is zero because recovery is byte exact. This example was selected to illustrate the GUI workflow and belongs to the training split, so it is not presented as held-out evidence.')
h('D. Software verification and reproducibility', 2)
p('The current regression suite passes 114 tests covering signal loading, prediction, password errors, tamper rejection, path boundaries, public previews, and exact restoration. Tests verify that a locked ciphertext preview needs neither model loading nor decryption and that wrong passwords do not return original data. Automated tests support functional correctness of these cases; they are not a cryptographic security proof, penetration test, or clinical validation.')
figure('fig04_camouflage_recovery.png', 'Fig. 4. Illustrative training fragment 100m (2).mat: original NSR score 99.75%, public-decoy APB score 94.64%, and exact restored original. Waveforms use raw dataset units; the difference panel magnifies the alteration.')
p('Classification figures originate from the saved results.json, split manifest, and validation predictions generated with the existing model. The audited camouflage batch is f5af518e8dc04ebb9bd4504b31c8f052. The checkpoint SHA-256 begins f2ea83714317c0b9; the full digest and batch configuration accompany the manuscript in the reproducibility notes. Reproduction uses the supplied MLII directory, checkpoint, current source code, and scripts. No model was retrained to create this manuscript.')

h('VI. Limitations and Responsible Use')
p('The strongest limitation is evaluation leakage. A new experiment should group fragments by verified patient or recording identity, remove duplicate waveforms before splitting, and report uncertainty across repeated independent splits. Rare classes require larger support. Current scores are not calibrated, and comparison with established ECG baselines on identical partitions remains outstanding. The results do not support a claim that the prototype diagnoses CVD or estimates individual patient risk.')
p('The threat model is narrow. Camouflage was optimized against one accessible checkpoint, without transfer testing, adaptive reconstruction attacks, membership inference, or clinical-attribute leakage analysis. The public waveform may retain clinically meaningful morphology even if the model changes its label. Fresh nonces and authenticated recovery protect the encrypted payload, but the public fixed password and retained plaintext source directory make the demo inappropriate for private patient records. A deployment would require separate key management, strong user secrets, access controls, and an explicit data-retention policy.')
p('The raw update uses an interpolated model-input gradient rather than an exact derivative through preprocessing. Its budget is relative to a per-record scale and is not a clinical distortion threshold. The FGSM curve samples only one attack family and cannot certify resistance to iterative, black-box, or adaptive attacks. No full RNAF reinforcement-learning policy or defense efficacy is evaluated. These unimplemented components should remain future research, not be described as completed contributions.')
p('The interface includes static, source-linked general care information for the original or restored model label. It never derives care advice from a decoy label, and the dataset contains no treatment or outcome targets. This feature is educational and has not undergone independent clinical validation; treatment selection requires qualified clinical review. Likewise, Grad-CAM and the uncalibrated model scores cannot validate a diagnosis.')

h('VII. Conclusion')
p('The workbench demonstrates a reproducible separation between adversarial ECG signal manipulation and authenticated exact-original recovery. A compact residual-attention baseline yields 87.00% accuracy on an internal fragment split, with substantial recording overlap that restricts the claim. The selected camouflage configuration changes all 1,000 recorded predictions, while independent checks recover every original MAT exactly. These findings establish the implemented workflow and its model sensitivity, not attack prevention, anonymization, or clinical readiness. Future evaluation should prioritize leakage-free recording or patient splits, stronger and transferable attacks, calibrated uncertainty, and secure key management before investigating a complete RNAF-inspired adaptive defense.')

h('References')
refs = [
    'P. Plawiak, “ECG signals (1000 fragments),” Mendeley Data, V3, 2017, doi: 10.17632/7dybx7wyfn.3.',
    'G. B. Moody and R. G. Mark, “The impact of the MIT-BIH Arrhythmia Database,” IEEE Eng. Med. Biol. Mag., vol. 20, no. 3, pp. 45-50, 2001. Database: PhysioNet, ver. 1.0.0, 2005, doi: 10.13026/C2F305.',
    'K. Rizwan, M. A. Habib, S. Raza, and M. Ahmad, “RNAF: ResilienceNet Adversarial Framework Using Deep Reinforcement Learning for Adversarial Attacks on Digital Images,” IEEE Access, vol. 13, pp. 201592-201610, 2025, doi: 10.1109/ACCESS.2025.3636942.',
    'I. J. Goodfellow, J. Shlens, and C. Szegedy, “Explaining and Harnessing Adversarial Examples,” in Proc. ICLR, 2015, arXiv:1412.6572.',
    'A. Madry, A. Makelov, L. Schmidt, D. Tsipras, and A. Vladu, “Towards Deep Learning Models Resistant to Adversarial Attacks,” in Proc. ICLR, 2018, arXiv:1706.06083.',
    'M. Dworkin, “Recommendation for Block Cipher Modes of Operation: Galois/Counter Mode (GCM) and GMAC,” NIST SP 800-38D, Nov. 2007, doi: 10.6028/NIST.SP.800-38D.',
    'R. R. Selvaraju, M. Cogswell, A. Das, R. Vedantam, D. Parikh, and D. Batra, “Grad-CAM: Visual Explanations from Deep Networks via Gradient-Based Localization,” Int. J. Comput. Vis., 2019, doi: 10.1007/s11263-019-01228-7.',
]
for i, text in enumerate(refs, 1):
    para = p(f'[{i}] {text}', indent=False, size=8)
    para.paragraph_format.left_indent = Inches(.18)
    para.paragraph_format.first_line_indent = Inches(-.18)
    para.paragraph_format.space_after = Pt(5)

# End the two-column section continuously so the final page is balanced.
end_section = doc.add_section(WD_SECTION_START.CONTINUOUS)
end_section._sectPr.find(qn('w:cols')).set(qn('w:num'), '1')
path = OUT / 'MITS_26MPE05015_IEEE_Manuscript.docx'
doc.save(path)
print(path)
